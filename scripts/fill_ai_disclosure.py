#!/usr/bin/env python3
"""
Populates the official Samsung PRISM Hackathon AI Usage Disclosure Form (LangAI3.0_AI_Disclosure.docx)
in accordance with approved disclosures and guidelines.
"""

from pathlib import Path
import docx
from docx.shared import Pt, Inches, RGBColor

DOCX_PATH = Path("/Users/manishharigovind/Desktop/VIT Vellore/PRISM/codedoctor/LangAI3.0_AI_Disclosure.docx")

def build_disclosure():
    doc = docx.Document(DOCX_PATH)
    
    # Clear existing template paragraphs while preserving document-level styles & margins
    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)

    def add_title(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(20)
        run.font.color.rgb = RGBColor(0x70, 0x50, 0x9B)  # Theme purple
        p.paragraph_format.space_after = Pt(12)
        return p

    def add_section_header(num, title, extra=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        r_num = p.add_run(f"{num}. ")
        r_num.bold = True
        r_num.font.name = "Calibri"
        r_num.font.size = Pt(13)
        r_title = p.add_run(title)
        r_title.bold = True
        r_title.font.name = "Calibri"
        r_title.font.size = Pt(13)
        r_title.font.color.rgb = RGBColor(0x30, 0x2C, 0x3C)
        if extra:
            r_ex = p.add_run(f" {extra}")
            r_ex.italic = True
            r_ex.font.name = "Calibri"
            r_ex.font.size = Pt(11)
            r_ex.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        return p

    def add_field(label, value):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r_lbl = p.add_run(f"{label}: ")
        r_lbl.bold = True
        r_lbl.font.name = "Calibri"
        r_lbl.font.size = Pt(11)
        r_val = p.add_run(value)
        r_val.font.name = "Calibri"
        r_val.font.size = Pt(11)
        return p

    def add_body(text, italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        r.italic = italic
        return p

    # --- Title
    add_title("AI Usage DISCLOSURE FORM")

    # --- 1. Team Details
    add_section_header("1", "Team Details")
    add_field("Team Name", "Doctors")
    add_field("Project / Product Name", "CodeDoctor — Agentic Semantic and Structural Intelligence for Code Repositories")
    add_field("Organization / Institution (if any)", "Vellore Institute of Technology (VIT), Vellore")
    add_field("Submission Date", "September 30, 2026")

    # --- 2. AI Usage Declaration
    add_section_header("2", "AI Usage Declaration")
    add_field("Did your team use any Artificial Intelligence (AI) in developing this project?", "Yes")
    add_body(
        "Declaration Summary: AI tools (Antigravity IDE, Claude 3.5 Sonnet, OpenAI Codex) were utilized as interactive "
        "pair-programming assistants during project development. AI assisted with architectural ideation, boilerplate code scaffolding, "
        "Tree-sitter AST visitor logic, unit test fixture generation, and technical documentation drafting. All core architectural gating, "
        "mathematical scoring formulas (Reciprocal Rank Fusion k=60, escalation margin Δ=0.004), content-addressed SQLite cache schemas, "
        "and experimental validation gates were designed, evaluated, and verified by the team."
    )

    # --- 3. Purpose of AI Usage (Brief Details)
    add_section_header("3", "Purpose of AI Usage (Brief Details)")
    purposes = [
        ("Idea generation / brainstorming", "Explored multi-signal code retrieval architectures, evidence-gated query routing mechanisms, and AST statement-boundary chunking strategies to minimize CPU query latency under Samsung constraints."),
        ("Code generation or assistance", "Assisted in writing boilerplate scaffolding, Tree-sitter AST visitor patterns, SQLite content-addressed caching methods (EmbeddingCache), and pytest test fixtures."),
        ("UI / UX design", "Assisted in structuring the local developer web workbench (web/dist/index.html and explorer.html), SVG icon alignments, and dark-mode CSS styling."),
        ("Content creation", "Assisted in drafting technical documentation (README.md, PROTOCOL.md), slide bullet points, and AI disclosure statements."),
        ("Data analysis", "Assisted in analyzing MTEB evaluation logs, computing paired-bootstrap statistical significance distributions, and evaluating Kaggle Intel Xeon CPU runtime profiles."),
        ("Testing / debugging", "Assisted in debugging test fixture assertions (e.g., scripts/validate_hands_on.py line 136 index-overwrite bug) and diagnosing cross-platform Windows/macOS subprocess isolation edge cases."),
        ("Other", "Assisted in authoring Kaggle multi-stage evaluation kernel scripts and GitHub Actions CI matrix workflow configurations.")
    ]
    for lbl, desc in purposes:
        add_field(lbl, desc)

    # --- 4. Feature Origin Classification
    add_section_header("4", "Feature Origin Classification", "(Please add/delete based on the number of the feature)")
    features = [
        (
            "Evidence-Gated Agentic Query Router (app/agent/router.py)",
            "Both",
            "Antigravity / Claude 3.5 Sonnet / Codex",
            "Design a query router that runs free passes (inverted index, BM25) first, defers expensive dense models on symbol queries, and escalates to dense semantic search only when the RRF fused score margin between rank 1 and 2 is below 0.004.",
            "Generated the initial routing state machine, RouterSignals dataclass, regex patterns for structural intent (who calls, what imports), and RetrievalPlan.",
            "Hand-tuned the score margin threshold (Δ = 0.004), integrated microsecond worked-example presence checks, and added audit-ready decision reason logging."
        ),
        (
            "Invariant Recipe-Keyed Vector Cache (app/indexing/cache.py)",
            "Both",
            "Antigravity / Claude 3.5 Sonnet",
            "Create a content-addressed SQLite embedding cache that keys on SHA-256 of model, commit revision, prompt, and chunk text while excluding file path and line numbers so moved functions don't re-embed.",
            "Provided SQLite table schema, binary vector serialization/deserialization routines, and hash computation functions.",
            "Added strict mode='document' vs mode='query' namespace isolation to prevent query-doc cache collisions, and verified bit-compatibility between GPU indexing and CPU retrieval (cosine >= 0.9994)."
        ),
        (
            "Tree-sitter AST Statement-Boundary Chunker (app/parser/chunker.py, edges.py)",
            "Both",
            "Antigravity / Codex",
            "Implement Tree-sitter chunking for Python and JavaScript/TypeScript that splits oversized functions strictly at top-level statement boundaries rather than arbitrary character/token cutoffs.",
            "Generated AST grammar bindings, declaration header signature extractors, and _split_oversized_function statement traversal logic.",
            "Enforced strict boundary conditions for Python def/class headers and JS arrow functions; added StructuralExtractor for call and import edge resolution with VERIFIED vs CANDIDATE confidence tags."
        ),
        (
            "Defensive Behavioral Execution Sandbox (app/pipeline/execution_rerank.py)",
            "Both",
            "Antigravity / Claude 3.5 Sonnet",
            "Implement an execution-based reranker that parses worked examples from competitive programming problems and runs candidate code in an isolated subprocess with timeout and memory limits.",
            "Produced regex parsers for Codeforces/AtCoder I/O blocks, temporary directory runner script generation, and subprocess timeout handling.",
            "Safety overhaul: moved setrlimit inside the child runner after exec to eliminate PyTorch multi-threading fork deadlocks and macOS kernel crashes; added floating-point tolerance (1e-06) and whitespace-invariant token matching (+0.02802 NDCG@10 lift)."
        ),
        (
            "Lineage-Based Evolutionary Disambiguation (app/indexing/lineage.py, versioned.py)",
            "Both",
            "Antigravity / Claude 3.5 Sonnet",
            "Design a multi-version indexer that clusters evolving function variants across Git commit history and ranks them using execution verification against test cases.",
            "Generated Git commit history traversal, k=5 token shingle hashing, and Jaccard similarity clustering algorithms.",
            "Integrated execution-first ranking logic within lineages, achieving 100% correct variant selection across 12 test lineages on Kaggle Xeon CPU (vs 8.33% for raw dense retrieval)."
        ),
        (
            "Local Repository Workbench & FastAPI Service (app/workbench.py, app/server.py)",
            "Both",
            "Antigravity / Claude 3.5 Sonnet",
            "Create a FastAPI backend and workbench orchestrator that serves multi-signal code search, exposes transparent additive RRF arithmetic, and displays retrieval decision reasons.",
            "Built FastAPI endpoint routes, Workbench orchestration class, and JSON schema response models.",
            "Enforced the zero-vector attention cap for oversized chunks (>6000 chars) to prevent corrupt interior-token pooling in last-token models; implemented test-split access budgeting to prevent data contamination."
        )
    ]

    for idx, (fname, origin, tools, prompt, out_summary, mod) in enumerate(features, 1):
        p_f = doc.add_paragraph()
        p_f.paragraph_format.space_before = Pt(8)
        p_f.paragraph_format.space_after = Pt(2)
        r_f = p_f.add_run(f"Feature {idx}: {fname}")
        r_f.bold = True
        r_f.font.name = "Calibri"
        r_f.font.size = Pt(11)
        r_f.font.color.rgb = RGBColor(0x70, 0x50, 0x9B)

        add_field("  1. Feature Name", fname)
        add_field("  2. Self-Generated/AI-Generated/Both", origin)
        desc = (
            f"AI Tools Used: {tools} | Prompt: \"{prompt}\" | "
            f"Output Summary: {out_summary} | Human Modification: {mod}"
        )
        add_field("  3. Description", desc)

    # --- 5. Ethical & Compliance Confirmation
    add_section_header("5", "Ethical & Compliance Confirmation")
    add_field("AI usage complies with guidelines and policies", "Yes")
    add_field("No proprietary or copyrighted data misused", "I Agree")

    # --- 6. Declaration & Sign-Off
    add_section_header("6", "Declaration & Sign-Off")
    add_field("Name of Team Representative", "Manish H.")
    add_field("Role", "Team Lead / Full-Stack Developer")
    add_field("Signature", "Manish H. (Digitally Signed)")
    add_field("Date", "September 30, 2026")

    doc.save(DOCX_PATH)
    print(f"Successfully populated and saved: {DOCX_PATH}")

if __name__ == "__main__":
    build_disclosure()
